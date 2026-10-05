import json
from datetime import timedelta
from io import BytesIO
import numpy as np
import pandas as pd
import pytest
from sqlalchemy import select
from src.core.db import session
from src.core.models import Product,Approval,Membership
from src.services.inventory import products,update_product,replenishment,propose,decide,execute
from src.services.data import sales,datasets
from src.services.analytics import forecast_product,finance
from src.analytics.forecasting import run_forecast,daily_series,predict
from src.analytics.customers import rfm,basket_rules,churn,segment_customers
from src.services.reports import build_report,export,save_report,get_report
from src.services.jobs import queue,claim,execute_job,get_job,cancel,schedule,tick

def test_forecast_holdout_is_not_used_for_selection():
    series=pd.Series(10+np.sin(np.arange(120)*2*np.pi/7),index=pd.date_range('2026-01-01',periods=120))
    changed=series.copy();changed.iloc[-14:]+=300
    a=run_forecast(series);b=run_forecast(changed)
    assert a['validation']==b['validation'] and a['model']==b['model']
    assert a['holdout']!=b['holdout']
    assert len(a['forecasts'])==14
    assert all(0<=x['lower']<=x['prediction']<=x['upper'] for x in a['forecasts'])

@pytest.mark.parametrize('method',['Seasonal naive','Moving average','Exponential smoothing','ARIMA','Croston SBA'])
def test_baseline_forecasts_finite(method):
    out=predict(np.array([0,1,2,3,0,0,4]*12),7,method)
    assert out.shape==(7,) and np.isfinite(out).all() and (out>=0).all()

def test_missing_dates_are_not_silently_zeros():
    frame=pd.DataFrame({'sku':['a','a'],'date':pd.to_datetime(['2026-01-01','2026-01-03']),'quantity':[2,4]})
    with pytest.raises(ValueError,match='Missing dates'):daily_series(frame,'a')
    assert daily_series(frame,'a',True).tolist()==[2,0,4]
    with pytest.raises(ValueError,match='28'):run_forecast(daily_series(frame,'a',True))

def test_inventory_budget_approval_receipt_once(demo,other):
    ctx,_=demo
    plan=replenishment(ctx,budget=12000)
    assert plan.estimated_cost.sum()<=12000
    sku='MILK-1';p=next(p for p in products(ctx) if p.sku==sku)
    aid=propose(ctx,'purchase','Buy milk',{'sku':sku,'quantity':12})
    with pytest.raises(ValueError):execute(ctx,aid)
    with pytest.raises(ValueError):decide(other,aid,True)
    decide(ctx,aid,True)
    with session() as s:s.get(Approval,aid).expires_at-=timedelta(days=30)
    assert replenishment(ctx).set_index('sku').loc[sku,'on_order']==12
    execute(ctx,aid)
    assert next(x for x in products(ctx) if x.sku==sku).on_hand==p.on_hand+12
    with pytest.raises(ValueError):execute(ctx,aid)
    assert replenishment(ctx).set_index('sku').loc[sku,'on_order']==0

def test_price_cannot_bypass_margin_or_changed_inputs(demo):
    ctx,_=demo
    with pytest.raises(ValueError):propose(ctx,'price','Too cheap',{'sku':'MILK-1','price':150})
    aid=propose(ctx,'price','Review price',{'sku':'MILK-1','price':320})
    update_product(ctx,'MILK-1',{'unit_cost':250})
    with pytest.raises(ValueError,match='changed'):decide(ctx,aid,True)

def test_analytics_and_jobs_end_to_end(demo):
    ctx,_=demo;d=sales(ctx)
    segments,metrics=segment_customers(d)
    assert len(segments)>=100 and metrics['k']>=2
    rules,n=basket_rules(d);assert n>=20 and len(rules)>0
    result=churn(d)
    if result['available']:assert 0<=result['pr_auc']<=1 and result['train_cutoff']<result['test_cutoff']
    else:assert result['reason']
    assert finance(ctx)['gross_profit']>0
    payload={'dataset_id':datasets(ctx)[0].id,'sku':'MILK-1','horizon':7,'method':'Seasonal naive'}
    jid=queue(ctx,'forecast',payload);assert queue(ctx,'forecast',payload)==jid
    assert claim()==jid and claim() is None
    execute_job(jid);job=get_job(ctx,jid)
    assert job.state=='completed' and len(job.result['forecasts'])==7
    rid=save_report(ctx,build_report('Forecast','Calculated result',{'Forecast':job.result['forecasts']},scope='sales'))
    assert get_report(ctx,rid)['tables']['Forecast'][0]['prediction']>=0

def test_cancel_and_schedule_idempotency(owner):
    sid=schedule(owner,'research',{'query':'public biscuits'},24)
    tick();first=claim();assert first
    tick();assert claim() is None
    cancel(owner,first);execute_job(first)
    assert get_job(owner,first).state=='cancelled'

def test_worker_runs_serialized_job_in_another_process(owner):
    import subprocess,sys
    from src.core.config import ROOT
    jid=queue(owner,'chat',{'question':'Explain the store policy','mode':'Evidence only'})
    assert claim()==jid
    run=subprocess.run([sys.executable,'-m','src.services.jobs','--run',jid],cwd=ROOT,capture_output=True,timeout=30)
    assert run.returncode==0,run.stderr.decode()
    job=get_job(owner,jid)
    assert job.state=='completed' and job.result['provider']=='Evidence tools'

def test_inventory_does_not_treat_missing_days_as_observed(owner):
    from src.services.data import normalize,suggest_mapping,import_sales
    frame=pd.DataFrame({'date':['2026-01-01','2026-01-03'],'sku':['A','A'],'quantity':[30,30],'price':[100,100],'unit_cost':[50,50]})
    clean,_,quality=normalize(frame,suggest_mapping(frame))
    import_sales(owner,'sparse.csv',frame.to_csv(index=False).encode(),clean,quality,False)
    row=replenishment(owner).iloc[0]
    assert row.suggested_quantity==0 and 'withheld' in row.source

@pytest.mark.parametrize('fmt',['txt','md','csv','xlsx','docx','pdf','json'])
def test_exports_preserve_numbers_and_neutralize_formulas(fmt):
    payload=build_report('Inventory','Source-backed totals',{'Rows':[{'sku':'=HYPERLINK("https://bad.test")','units':12.5,'missing':float('nan')}]},['Dataset D1'],['Review before ordering'],'inventory')
    result=export(payload,fmt);assert len(result)>80
    if fmt=='xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(result));assert wb['Rows']['B2'].value==12.5
        assert wb['Rows']['A2'].data_type=='s' and wb['Rows']['A2'].value.startswith("'")
        assert wb['Rows']['C2'].value is None
    if fmt=='docx':
        from docx import Document
        d=Document(BytesIO(result));assert d.tables[0].cell(1,1).text=='12.5'
    if fmt=='pdf':
        from pypdf import PdfReader
        assert '12.5' in ''.join(p.extract_text() for p in PdfReader(BytesIO(result)).pages)
    if fmt=='csv':assert "'=HYPERLINK" in result.decode('utf-8-sig')
    if fmt=='json':assert json.loads(result)['tables']['Rows'][0]['missing'] is None

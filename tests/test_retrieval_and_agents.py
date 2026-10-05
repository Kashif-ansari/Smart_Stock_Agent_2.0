from pathlib import Path
from io import BytesIO
import socket,json
import numpy as np
import pytest
from src.rag import retrieval
from src.rag.extract import extract,chunk_segments
from src.core.config import tenant_dir
from src.agents.pipeline import run_case,check_answer
from src.services.market import safe_url,assess_opportunity

class Embedder:
    def passage_embed(self,texts):return [np.array([1.,.1]) if 'milk' in x.lower() else np.array([.1,1.]) for x in texts]
    def query_embed(self,q):return self.passage_embed([q])

class Reranker:
    def rerank(self,q,texts):return [.9 if 'milk' in t.lower() else .1 for t in texts]

def test_faiss_hybrid_reranking_integrity_and_stale_index(owner,monkeypatch):
    monkeypatch.setenv('ENABLE_SEMANTIC_SEARCH','true')
    monkeypatch.setattr(retrieval,'embedding_model',lambda _:Embedder())
    monkeypatch.setattr(retrieval,'reranker_model',lambda _:Reranker())
    retrieval.add_document(owner,'milk.txt',b'Milk SKU-1 must be refrigerated.')
    build=retrieval.build_index(owner)
    hits,meta=retrieval.retrieve(owner,'milk SKU-1')
    assert hits and hits[0]['title']=='milk.txt' and 'cross-encoder' in meta['mode']
    root=tenant_dir(owner.tenant_id,'indexes')/'general'/build['version']
    (root/'index.faiss').write_bytes(b'not a trusted index')
    hits,meta=retrieval.retrieve(owner,'milk');assert hits and any('integrity' in x for x in meta['notes'])
    retrieval.add_document(owner,'more.txt',b'More milk receiving instructions')
    _,meta=retrieval.retrieve(owner,'milk');assert any('rebuild' in x for x in meta['notes'])

def test_extraction_preserves_docx_order_and_chunk_limits():
    from docx import Document
    d=Document();d.add_heading('Receiving',1);d.add_paragraph('First paragraph')
    t=d.add_table(rows=2,cols=2);t.cell(0,0).text='SKU';t.cell(0,1).text='Lead days';t.cell(1,0).text='MILK-1';t.cell(1,1).text='3'
    d.add_paragraph('Last paragraph');b=BytesIO();d.save(b)
    parts=extract('policy.docx',b.getvalue())
    assert parts[-1]['text']=='Last paragraph'
    assert 'SKU: MILK-1' in parts[-2]['text'] and 'Table' in parts[-2]['location']
    chunks=chunk_segments([{'text':' '.join(['word']*500),'location':'Page 1'},{'text':' '.join(['text']*210),'location':'Page 2'}])
    assert all(len(c['text'].split())<=220 for c in chunks)

@pytest.mark.parametrize('url',['http://127.0.0.1/','http://169.254.169.254/','file:///etc/passwd','http://user:pass@example.com/','http://example.com:8000/'])
def test_ssrf_blocks_nonpublic_urls(url):
    with pytest.raises((ValueError,OSError)):safe_url(url)

def test_ssrf_rejects_mixed_dns(monkeypatch):
    monkeypatch.setattr(socket,'getaddrinfo',lambda *a,**kw:[(2,1,6,'',('8.8.8.8',443)),(2,1,6,'',('127.0.0.1',443))])
    with pytest.raises(ValueError):safe_url('https://example.test/')

def test_evidence_pipeline_and_unsupported_sources(demo,monkeypatch):
    ctx,_=demo
    result=run_case(ctx,'Which inventory products should I reorder?')
    assert result['tables']['Replenishment'] and result['citations']
    assert [s['stage'] for s in result['stages']][-1]=='Final report'
    assert check_answer({'answer':'Sales are 3 [S999]','citations':['S999']},[{'id':'S1'}])
    assert check_answer({'answer':'Fine','citations':None},[])
    with pytest.raises(ValueError,match='requires|requires|Live AI'):run_case(ctx,'sales revenue','Groq assistant')

def test_missing_evidence_does_not_invent_numbers(owner):
    result=run_case(owner,'How many units will I sell next week?')
    assert 'not have enough' in result['answer'] and not result['tables']

def test_market_catalog_match_separates_stockout(demo):
    ctx,_=demo
    from src.services.inventory import update_product
    update_product(ctx,'MILK-1',{'on_hand':0})
    report=assess_opportunity(ctx,'Whole milk 1L',231,320,12,5000,True,7,[])
    assert 'out of stock' in report['tables']['Pilot scenario'][0]['catalog_status']
    assert report['summary'].startswith('Hold')

@pytest.mark.optional
def test_crewai_real_flow_with_mocked_provider(monkeypatch):
    pytest.importorskip('crewai')
    from src.agents import crew
    def fake(messages):
        text=str(messages).lower()
        if 'evidence reviewer' in text or 'supported (boolean)' in text:return json.dumps({'supported':True,'issues':[]})
        return json.dumps({'answer':'Milk requires refrigeration [S1].','citations':['S1']})
    monkeypatch.setattr(crew,'groq_complete',fake)
    answer,review=crew.run_crew([{'role':'user','content':'Explain milk storage [S1]'}],[{'id':'S1','text':'Milk requires refrigeration.'}])
    assert answer['citations']==['S1'] and 'CrewAI' in review
